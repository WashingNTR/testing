import streamlit as st
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from tempfile import NamedTemporaryFile
from io import BytesIO

st.set_page_config(page_title="Transformation Excel", layout="centered")
st.title("📊 Application de transformation Excel (sans VBA)")

uploaded_file = st.file_uploader("Déposez un fichier Excel (.xlsx)", type=["xlsx"])

def traiter_fichier(fichier):
    mot_de_passe = "newrest2025"
    wb = load_workbook(fichier)

    if "Export" not in wb.sheetnames:
        st.error("La feuille 'Export' est manquante.")
        return None

    ws_export = wb["Export"]

    # Supprimer toutes les autres feuilles
    for sheet in wb.sheetnames[:]:
        if sheet != "Export":
            del wb[sheet]

    # Supprimer les colonnes J et suivantes
    last_col = ws_export.max_column
    for col in reversed(range(10, last_col + 1)):
        ws_export.delete_cols(col)

    # Ajouter les colonnes J à N
    ws_export.insert_cols(10, amount=5)
    headers = ["Accepté/Refusé", "Commentaire", "Autre", "KAM / TO", "Commentaire"]
    for i, header in enumerate(headers):
        ws_export.cell(row=1, column=10 + i, value=header)

    # Validation de données dans colonnes J et M
    val_list = DataValidation(type="list", formula1='"Accepté,Refusé,N/A"', allow_blank=True)
    ws_export.add_data_validation(val_list)
    val_list.add(f"J2:J{ws_export.max_row}")

    val_list2 = DataValidation(type="list", formula1='"Accepté,Refusé,N/A"', allow_blank=True)
    ws_export.add_data_validation(val_list2)
    val_list2.add(f"M2:M{ws_export.max_row}")

    # Format de date sur la colonne A
    for row in ws_export.iter_rows(min_row=2, min_col=1, max_col=1):
        for cell in row:
            cell.number_format = "dddd dd mmmm yyyy"

    # Wrap text et largeur
    wrap_cols = ["F", "G", "H", "I", "J", "K", "M", "N"]
    for col in wrap_cols:
        ws_export.column_dimensions[col].width = 30
    for row in ws_export.iter_rows(min_row=2, max_col=14):
        for cell in row:
            if cell.column_letter in wrap_cols:
                cell.alignment = Alignment(wrapText=True)

    # Protection des cellules
    for row in ws_export.iter_rows(min_row=2, max_row=ws_export.max_row, max_col=14):
        for cell in row:
            if cell.column <= 9 or cell.column in [13, 14]:  # A:I, M:N
                cell.protection = Protection(locked=True)
            else:
                cell.protection = Protection(locked=False)

    ws_export.protection.sheet = True
    ws_export.protection.password = mot_de_passe

    # Création des feuilles par site
    sites_specifiques = {"ORY", "MRS", "LYS", "NTE", "BRU", "MPL", "RNS", "BOD", "TLS"}
    feuilles = {}

    for row in ws_export.iter_rows(min_row=2, max_row=ws_export.max_row):
        site = row[3].value
        if site:
            nom_feuille = site if site in sites_specifiques else "Autre"
            if nom_feuille not in feuilles:
                feuilles[nom_feuille] = wb.create_sheet(title=nom_feuille)
                for i, cell in enumerate(ws_export[1], start=1):
                    feuilles[nom_feuille].cell(row=1, column=i).value = cell.value
            feuilles[nom_feuille].append([cell.value for cell in row])

    # Création de la feuille Consolidation
    if "Consolidation" in wb.sheetnames:
        del wb["Consolidation"]
    ws_consol = wb.create_sheet(title="Consolidation")
    ws_consol.append([cell.value for cell in ws_export[1]])

    for feuille, ws in feuilles.items():
        for row in ws.iter_rows(min_row=2, values_only=True):
            ws_consol.append(row)

    # Validation dans la colonne M de la consolidation
    val_list3 = DataValidation(type="list", formula1='"Accepté,Refusé,N/A"', allow_blank=True)
    ws_consol.add_data_validation(val_list3)
    val_list3.add(f"M2:M{ws_consol.max_row}")

    # Formatage de la consolidation
    for col in ["J", "K", "M", "N"]:
        ws_consol.column_dimensions[col].width = 30
    for row in ws_consol.iter_rows(min_row=2, max_col=14):
        for cell in row:
            if cell.column_letter in wrap_cols:
                cell.alignment = Alignment(wrapText=True)

    # Protection des cellules de Consolidation
    for row in ws_consol.iter_rows(min_row=2, max_row=ws_consol.max_row, max_col=14):
        for cell in row:
            if cell.column <= 9 or cell.column in [13, 14]:
                cell.protection = Protection(locked=True)
            else:
                cell.protection = Protection(locked=False)

    ws_consol.protection.sheet = True
    ws_consol.protection.password = mot_de_passe

    return wb


if uploaded_file:
    if st.button("Lancer le traitement"):
        with st.spinner("Traitement en cours..."):
            workbook = traiter_fichier(uploaded_file)

            if workbook:
                output = BytesIO()
                workbook.save(output)
                st.success("Traitement terminé avec succès ✅")
                st.download_button(
                    label="📥 Télécharger le fichier transformé",
                    data=output.getvalue(),
                    file_name="fichier_traite.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
